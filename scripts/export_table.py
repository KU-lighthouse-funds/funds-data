#!/usr/bin/env python3
"""Export cleaned CSV to Excel and a local HTML table viewer."""
import csv
import html
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(r"c:\Users\Kaja\Documents\Funds")
CSV_PATH = BASE / "updated file without VC - cleaned.csv"
XLSX_PATH = BASE / "updated file without VC - cleaned.xlsx"
HTML_PATH = BASE / "funds-table.html"


def load_rows():
    with CSV_PATH.open(encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))
    return rows[0], rows[1:]


def export_xlsx(headers, data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Funds"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")

    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for r, row in enumerate(data, 2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if headers[c - 1] == "Link" and val.startswith("http"):
                cell.hyperlink = val
                cell.font = Font(color="0563C1", underline="single")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(data) + 1}"
    widths = [18, 28, 35, 40, 22, 18, 10, 18, 16, 55, 18, 22, 24]
    for i, w in enumerate(widths[: len(headers)], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(XLSX_PATH)


def export_html(headers, data):
    header_html = "".join(f"<th>{html.escape(h)}</th>" for h in headers)
    body_rows = []
    for row in data:
        cells = []
        for i, val in enumerate(row):
            if headers[i] == "Link" and val.startswith("http"):
                cells.append(
                    f'<td><a href="{html.escape(val)}" target="_blank" rel="noopener">Open</a></td>'
                )
            else:
                cells.append(f"<td>{html.escape(val)}</td>")
        body_rows.append("<tr>" + "".join(cells) + "</tr>")

    page = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Funds table</title>
<style>
  * {{ box-sizing: border-box; }}
  body {{ font-family: Segoe UI, system-ui, sans-serif; margin: 0; padding: 1rem 1.5rem; background: #f5f7fa; color: #1a1a1a; }}
  h1 {{ margin: 0 0 .25rem; font-size: 1.5rem; }}
  p {{ margin: 0 0 1rem; color: #555; }}
  .toolbar {{ display: flex; gap: .75rem; flex-wrap: wrap; margin-bottom: 1rem; align-items: center; }}
  #search {{ padding: .5rem .75rem; border: 1px solid #ccc; border-radius: 6px; min-width: 260px; font-size: 1rem; }}
  .count {{ color: #666; font-size: .9rem; }}
  .wrap {{ overflow: auto; background: #fff; border-radius: 8px; box-shadow: 0 1px 4px rgba(0,0,0,.08); max-height: calc(100vh - 140px); }}
  table {{ border-collapse: collapse; width: max-content; min-width: 100%; font-size: .85rem; }}
  th, td {{ border-bottom: 1px solid #e8e8e8; padding: .5rem .6rem; text-align: left; vertical-align: top; max-width: 320px; }}
  th {{ position: sticky; top: 0; background: #1F4E79; color: #fff; z-index: 1; white-space: nowrap; }}
  tr:hover td {{ background: #f0f6ff; }}
  a {{ color: #0563C1; }}
</style>
</head>
<body>
<h1>Funds and programmes</h1>
<p>Local table view. Open this file in any browser and use search to filter rows.</p>
<div class="toolbar">
  <input type="search" id="search" placeholder="Search name, segment, stage, CVR..." autofocus>
  <span class="count" id="count"></span>
</div>
<div class="wrap">
<table id="tbl">
<thead><tr>{header_html}</tr></thead>
<tbody>
{"".join(body_rows)}
</tbody>
</table>
</div>
<script>
const search = document.getElementById('search');
const rows = [...document.querySelectorAll('#tbl tbody tr')];
const count = document.getElementById('count');
function filter() {{
  const q = search.value.toLowerCase().trim();
  let n = 0;
  rows.forEach(r => {{
    const show = !q || r.textContent.toLowerCase().includes(q);
    r.style.display = show ? '' : 'none';
    if (show) n++;
  }});
  count.textContent = n + ' / ' + rows.length + ' programmes';
}}
search.addEventListener('input', filter);
filter();
</script>
</body>
</html>"""
    HTML_PATH.write_text(page, encoding="utf-8")


def main():
    headers, data = load_rows()
    export_xlsx(headers, data)
    export_html(headers, data)
    print(f"Excel: {XLSX_PATH}")
    print(f"HTML:  {HTML_PATH}")


if __name__ == "__main__":
    main()
