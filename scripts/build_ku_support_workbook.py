"""
Build a NEW funds workbook that keeps the existing table structure
and adds KU Preaward / Lighthouse support columns from the KU LH PPT.

Does NOT modify: updated file without VC - cleaned.csv / .xlsx
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(r"c:\Users\Kaja\Documents\Funds")
SRC_CSV = BASE / "updated file without VC - cleaned.csv"
OUT_CSV = BASE / "funds with KU support - v1.csv"
OUT_XLSX = BASE / "funds with KU support - v1.xlsx"

# New columns appended (structure of existing columns unchanged)
NEW_COLS = [
    "KU support unit",  # Preaward | Lighthouse | Joint | —
    "KU faculty focus",  # SCIENCE, SUND, Alle, …
    "KU contact hint",  # who researchers should reach first
    "PPT notes",  # extra facts from PPT (does not replace Quick info)
]

# Default contact routing (from PPT slide 8)
DEFAULT_CONTACTS = {
    "Preaward": (
        "Preaward RSO — Frederiksberg+ (SCIENCE): Agnieszka Zygadlo Nielsen; "
        "Nørre (SUND): Peter Lundberg & Esben Fiedler Røge; Søndre (SSH): TBD. "
        "If in doubt: Julie Sand Jørgensen / KU Lighthouse team."
    ),
    "Lighthouse": (
        "KU Lighthouse Innovation Advisors (see KU Lighthouse team page). "
        "If in doubt: Julie Sand Jørgensen."
    ),
    "Joint": (
        "Case-centric team with Preaward + Lighthouse. "
        "Start with the unit listed as primary below, or Julie Sand Jørgensen if unsure."
    ),
}

# Explicit mappings: our table Name -> support metadata from PPT
# Sources: slides 11 (joint contact), 13–14 (Preaward), 15–16 (Lighthouse), detail slides
SUPPORT: dict[str, dict] = {
    # --- NNF ---
    "NNF Pioneer Innovator Grant": {
        "unit": "Preaward",
        "faculty": "SCIENCE, SUND",
        "contact": "Preaward RSO (primary). NNF: Signe Daugbjerg (Health) / Charlotte Schöller (Sustainability).",
        "notes": "PPT: up to DKK 1.2M / 1 year; TRL 1–3; pre-CVR PoC. Health or Sustainability themes.",
    },
    "NNF Distinguished Innovator Grant": {
        "unit": "Preaward",
        "faculty": "SCIENCE, SUND",
        "contact": "Preaward RSO (primary). NNF: Signe Daugbjerg (Health) / Charlotte Schöller (Sustainability).",
        "notes": "PPT: up to DKK 7.6M / 3 years; TRL 3–5; pre-CVR; senior innovators.",
    },
    "Open Discovery Innovation Network (ODIN)": {
        "unit": "Preaward",
        "faculty": "SCIENCE, SUND (health discovery)",
        "contact": "Preaward RSO (primary). ODIN Secretariat, Aarhus University.",
        "notes": "PPT: Open Innovation in Science; patent-free university–industry collaboration. DKK 235M platform (2024–2028).",
    },
    # --- IFD ---
    "IFD Innoexplorer": {
        "unit": "Lighthouse",
        "faculty": "Alle – kommercielt fokus",
        "contact": "KU Lighthouse (primary per PPT joint list). Fund inbox: Innoexplorer@innofond.dk",
        "notes": "PPT: DKK 500K–1.5M; up to 12 months; TRL 3–5; public researchers. Next round noted: 23 Sep 2026.",
    },
    "IFD Innofounder": {
        "unit": "Lighthouse",
        "faculty": "Alle – kommercielt fokus",
        "contact": "KU Lighthouse. Fund inbox: innofounder@innofond.dk",
        "notes": "PPT: DKK 27.5–82.5K/month (1–3 founders) + DKK 100K development; 12 months; TRL 1–3.",
    },
    "IFD Innobooster": {
        "unit": "Lighthouse",
        "faculty": "Alle – kommercielt fokus",
        "contact": "KU Lighthouse. Fund inbox: Innobooster@innofond.dk",
        "notes": "PPT: DKK 200K–5M; max 35% co-financing; up to 24 months; TRL 4–7; Danish SMEs/start-ups.",
    },
    "IFD Grand Solutions": {
        "unit": "Preaward",
        "faculty": "Alle",
        "contact": "Preaward RSO (primary). Expect FRB+/Nørre routing by faculty for case handling.",
        "notes": "PPT: DKK 5–40M; up to 5 years; 90% funding rate for Danish universities; TRL 3–8. Phase deadlines noted ~2027.",
    },
    "IFD Industrial Researcher (Industrial PhD/Postdoc)": {
        "unit": "Preaward",
        "faculty": "Alle",
        "contact": "Preaward RSO",
        "notes": "PPT lists ErhvervsPhD / ErhvervsPostdoc under Preaward-managed IFD calls.",
    },
    # --- Horizon / EIC ---
    "Horizon Europe": {
        "unit": "Preaward",
        "faculty": "Alle (cluster-dependent)",
        "contact": "Preaward RSO (clusters / collaborative projects)",
        "notes": "PPT: HE Clusters under Preaward. EUopSTART also Preaward.",
    },
    "EIC Pathfinder": {
        "unit": "Preaward",
        "faculty": "Alle",
        "contact": "Preaward RSO",
        "notes": "PPT joint list: EIC Pathfinder Open/Challenges → Preaward.",
    },
    "EIC Transition": {
        "unit": "Lighthouse",
        "faculty": "Alle – kommercielt fokus",
        "contact": "KU Lighthouse (primary per PPT joint list)",
        "notes": "PPT: EUR 0.5–2.5M (2026); 1–3 years; must build on eligible prior project results. Deadline noted 16 Sep 2026.",
    },
    "EIC Accelerator": {
        "unit": "Lighthouse",
        "faculty": "Alle – kommercielt fokus",
        "contact": "KU Lighthouse / EIC NCP",
        "notes": "PPT HE EIC Accelerator-style detail: high TRL scale-up; grant + equity. Confirm case routing with LH if researcher-linked.",
    },
    "Eurostars": {
        "unit": "Preaward",
        "faculty": "Alle",
        "contact": "Preaward RSO",
        "notes": "PPT joint list: Eurostar → Preaward.",
    },
    "Eureka": {
        "unit": "Preaward",
        "faculty": "Alle",
        "contact": "Preaward RSO",
        "notes": "Related international innovation network; treat like Eurostars routing unless told otherwise.",
    },
    # --- Demo / enviro funds ---
    "GUDP": {
        "unit": "Preaward",
        "faculty": "SCIENCE, SUND",
        "contact": "Preaward RSO (Frederiksberg+ / relevant faculty)",
        "notes": "PPT: demonstration/development projects under Preaward; SCIENCE, SUND.",
    },
    "MUDP": {
        "unit": "Preaward",
        "faculty": "SCIENCE",
        "contact": "Preaward RSO (Frederiksberg+)",
        "notes": "PPT: Preaward-managed; SCIENCE.",
    },
    "EUDP": {
        "unit": "Preaward",
        "faculty": "SCIENCE",
        "contact": "Preaward RSO",
        "notes": "Energy demo programme; route like other ministry demonstration schemes via Preaward unless LH advises otherwise.",
    },
    # --- UCPH internal / LH portfolio ---
    "UCPH Proof of Concept Fund (POC MAX)": {
        "unit": "Lighthouse",
        "faculty": "Alle (patentable inventions)",
        "contact": "POC@adm.ku.dk / KU Lighthouse",
        "notes": "PPT: up to DKK 500K; up to 18 months; 2–3 researchers + business partner; commercial only; TRL 3–9.",
    },
    "UCPH Proof of Concept Fund (POC-TO-GO)": {
        "unit": "Lighthouse",
        "faculty": "Alle",
        "contact": "POC@adm.ku.dk / KU Lighthouse",
        "notes": "PPT: up to DKK 150K; up to 6 months; 1–3 researchers + business partner; commercial or non-commercial; TRL 1–3.",
    },
    "Lighthouse Launch": {
        "unit": "Lighthouse",
        "faculty": "Alle (students)",
        "contact": "KU Lighthouse",
        "notes": "KU Lighthouse student pre-accelerator (not detailed as a separate PPT funding matrix row).",
    },
    "SPARK Denmark": {
        "unit": "Lighthouse",
        "faculty": "SUND, SCIENCE",
        "contact": "SPARK Denmark Secretariat (SUND) / KU Lighthouse",
        "notes": "PPT: up to DKK 700K/project; life-science translation mentoring + funding.",
    },
    "Spin-outs Denmark": {
        "unit": "Lighthouse",
        "faculty": "SUND, SCIENCE",
        "contact": "KU Lighthouse / tech transfer; Spin-outs Denmark secretariat",
        "notes": "PPT: VILLUM-funded national postdoc spin-out pathway; junior researchers.",
    },
    "Open Entrepreneurship": {
        "unit": "Lighthouse",
        "faculty": "Alle – kommercielt fokus",
        "contact": "KU Lighthouse / OE at UCPH; central unit DTU",
        "notes": "PPT: match researchers with external entrepreneurs/investors.",
    },
    "BII Bio Studio": {
        "unit": "Lighthouse",
        "faculty": "SUND, SCIENCE",
        "contact": "KU Lighthouse",
        "notes": "PPT: BII portfolio under Lighthouse-managed innovation programmes; company-creation stage.",
    },
    "BII Venture Lab": {
        "unit": "Lighthouse",
        "faculty": "SUND, SCIENCE",
        "contact": "KU Lighthouse",
        "notes": "PPT: early-stage life science & deep tech; ~€500K convertible loan / 12 months.",
    },
    "BII Quantum Lab": {
        "unit": "Lighthouse",
        "faculty": "SUND, SCIENCE",
        "contact": "KU Lighthouse",
        "notes": "PPT: BII Quantum Lab listed under Lighthouse-managed programmes.",
    },
    "Otto Mønsteds Fond – Den Lyse Ide": {
        "unit": "Lighthouse",
        "faculty": "Alle",
        "contact": "KU Lighthouse",
        "notes": "PPT: Den Lyse Idé under Lighthouse Otto Mønsted portfolio.",
    },
    "Otto Bruuns Fond": {
        "unit": "Lighthouse",
        "faculty": "Hard Tech / Alle (programme-dependent)",
        "contact": "KU Lighthouse",
        "notes": "PPT: industrielle projekter (Hard Tech) and almennyttige projekter (Alle).",
    },
    "Mikrolegat": {
        "unit": "Lighthouse",
        "faculty": "Alle – kommercielt fokus",
        "contact": "KU Lighthouse / Fonden for Entreprenørskab channels",
        "notes": "PPT: Mikrolegat + Momentum under Lighthouse-managed list.",
    },
    "Fonden for Entreprenørskab": {
        "unit": "Lighthouse",
        "faculty": "Alle – kommercielt fokus",
        "contact": "KU Lighthouse",
        "notes": "PPT: project support to schools/education institutions + Mikrolegat family.",
    },
    "Danmarks Entreprenørskabsfestival": {
        "unit": "Lighthouse",
        "faculty": "Alle – education",
        "contact": "KU Lighthouse / Fonden for Entreprenørskab",
        "notes": "Linked to Fonden for Entreprenørskab ecosystem (Lighthouse list).",
    },
    "Vissing Fonden": {
        "unit": "Lighthouse",
        "faculty": "SUND / Energi (programme-dependent)",
        "contact": "KU Lighthouse",
        "notes": "PPT: medical research (SUND); energy startups (Energi); social youth projects also exist.",
    },
    # Lundbeck
    "Lundbeck Frontier": {
        "unit": "Preaward",
        "faculty": "SCIENCE, SUND",
        "contact": "Preaward RSO",
        "notes": "PPT lists Lundbeck Start-Up Programme under Preaward (SCIENCE, SUND).",
    },
}


def load_src() -> tuple[list[str], list[dict]]:
    with SRC_CSV.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=";")
        rows = list(reader)
        return list(reader.fieldnames or []), rows


def enrich(rows: list[dict]) -> list[dict]:
    out = []
    for row in rows:
        name = row["Name"]
        meta = SUPPORT.get(name)
        if meta:
            unit = meta["unit"]
            faculty = meta.get("faculty", "")
            contact = meta.get("contact") or DEFAULT_CONTACTS.get(unit, "")
            notes = meta.get("notes", "")
        else:
            unit = "—"
            faculty = ""
            contact = ""
            notes = (
                "Not listed as a KU Preaward/Lighthouse-managed programme in the KU LH PPT "
                "(researcher help routing unknown / external)."
            )
        row = dict(row)
        row["KU support unit"] = unit
        row["KU faculty focus"] = faculty
        row["KU contact hint"] = contact
        row["PPT notes"] = notes
        out.append(row)
    return out


def write_csv(headers: list[str], rows: list[dict]) -> None:
    with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(headers: list[str], rows: list[dict]) -> None:
    wb = Workbook()

    # Sheet 1: Funds
    ws = wb.active
    ws.title = "Funds"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    ku_fill = PatternFill("solid", fgColor="2E7D32")
    header_font = Font(bold=True, color="FFFFFF")

    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = ku_fill if h in NEW_COLS else header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for r, row in enumerate(rows, 2):
        for c, h in enumerate(headers, 1):
            val = row.get(h, "")
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if h == "Link" and isinstance(val, str) and val.startswith("http"):
                cell.hyperlink = val
                cell.font = Font(color="0563C1", underline="single")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    widths = {
        "Opportunity": 18,
        "Name": 32,
        "Link": 34,
        "Criteria": 36,
        "Industrial segment": 22,
        "Stage": 16,
        "Geography": 12,
        "Funding Amount": 28,
        "Deadline": 22,
        "Quick info": 48,
        "CVR at application": 14,
        "CVR at programme start": 16,
        "KU support unit": 14,
        "KU faculty focus": 22,
        "KU contact hint": 40,
        "PPT notes": 50,
    }
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 18)

    # Sheet 2: Legend / how to use
    leg = wb.create_sheet("KU support legend")
    leg["A1"] = "KU Preaward / Lighthouse support — reading guide"
    leg["A1"].font = Font(bold=True, size=14)
    legend_rows = [
        "",
        "Source PPT: KU LH preaward funding support-AZ-JSJ.pptx",
        "Base table: updated file without VC - cleaned.csv (unchanged)",
        "",
        "Column meanings",
        "KU support unit — Which KU unit typically leads preaward help for this opportunity:",
        "  Preaward = Research Support Offices (faculty preaward teams)",
        "  Lighthouse = KU Lighthouse innovation advisors",
        "  Joint = both may be involved; see contact hint for primary",
        "  — = not in PPT KU-managed lists (external / no KU routing claimed)",
        "",
        "KU faculty focus — Faculty/audience hints from PPT matrices (SCIENCE, SUND, Alle, …).",
        "KU contact hint — Who researchers should reach first (PPT slide 8 + programme slides).",
        "PPT notes — Extra facts from the PPT; does NOT replace Quick info.",
        "",
        "Default Preaward contacts (PPT slide 8)",
        "Frederiksberg+ (SCIENCE): Agnieszka Zygadlo Nielsen",
        "Nørre (SUND): Peter Lundberg and Esben Fiedler Røge",
        "Søndre (SSH): TBD",
        "If in doubt: Julie Sand Jørgensen / KU Lighthouse team page",
        "",
        "What researchers can expect from Preaward (supporting role)",
        "1–2 professional readings of application per researcher per programme",
        "1–2 meetings of ~1h per application per researcher per programme",
        "",
        "Programmes in PPT but not yet as own rows in the funds table (candidates to add later)",
        "Plant2Food (NNF OIS) — Preaward",
        "IBIS (NNF biofertiliser OIS) — Preaward",
        "BII Upscalator / AI Lab / Venture House — Lighthouse (BII family)",
        "Plantefonden — Preaward (SCIENCE, SUND)",
        "EUopSTART — Preaward",
        "Global Innovation Network Programme (GINP) — Preaward",
        "Scaleup Europe Fund — Lighthouse list",
        "EIFO instruments — Lighthouse list (investment/guarantee; often VC-adjacent)",
        "Industriens Fond / DI fonde & legater — Lighthouse list",
    ]
    for i, text in enumerate(legend_rows, 1):
        leg.cell(row=i, column=1, value=text)
    leg.column_dimensions["A"].width = 110

    # Sheet 3: coverage summary
    cov = wb.create_sheet("KU coverage summary")
    cov.append(["KU support unit", "Count"])
    from collections import Counter

    counts = Counter(r["KU support unit"] for r in rows)
    for k, v in sorted(counts.items(), key=lambda x: (-x[1], x[0])):
        cov.append([k, v])
    cov.append([])
    cov.append(["Name", "KU support unit", "KU faculty focus"])
    for r in sorted(rows, key=lambda x: (x["KU support unit"], x["Name"])):
        if r["KU support unit"] != "—":
            cov.append([r["Name"], r["KU support unit"], r["KU faculty focus"]])

    wb.save(OUT_XLSX)


def main() -> None:
    base_headers, rows = load_src()
    headers = base_headers + [c for c in NEW_COLS if c not in base_headers]
    enriched = enrich(rows)
    write_csv(headers, enriched)
    write_xlsx(headers, enriched)

    mapped = sum(1 for r in enriched if r["KU support unit"] != "—")
    print(f"Source rows: {len(rows)}")
    print(f"Mapped to KU support: {mapped}")
    print(f"Unmapped (—): {len(rows) - mapped}")
    print(f"CSV:  {OUT_CSV}")
    print(f"XLSX: {OUT_XLSX}")


if __name__ == "__main__":
    main()
